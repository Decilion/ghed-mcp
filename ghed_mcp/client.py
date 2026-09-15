"""Workbook download, cache, and provenance helpers."""
from __future__ import annotations

import asyncio
from contextlib import asynccontextmanager
import tempfile
import time
import zipfile
import logging
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
from filelock import FileLock, Timeout
from openpyxl import load_workbook

logger = logging.getLogger("ghed_mcp.client")

BASE_URL = "https://apps.who.int"
DOCUMENTATION_TREE_URL = (
    "https://apps.who.int/nha/database/DocumentationCentre/GetTree/en"
)
LEGACY_SOURCE_URL = "https://apps.who.int/nha/database/Home/IndicatorsDownload/en"
USER_AGENT = "mcp-server-ghed/0.6.0 (+https://decilion.com)"
DEFAULT_TIMEOUT = 120.0
_MAX_DOWNLOAD_ATTEMPTS = 3
_BACKOFF_BASE_SECONDS = 2.0


def _is_retryable_status(status: int) -> bool:
    return status >= 500 or status == 429


class GHEDError(Exception):
    """Structured error from GHED download or workbook handling."""

    def __init__(
        self,
        message: str,
        *,
        status: int | None = None,
        source_url: str = DOCUMENTATION_TREE_URL,
        path: str | None = None,
    ):
        super().__init__(message)
        self.status = status
        self.source_url = source_url
        self.path = path

    def to_dict(self) -> dict[str, Any]:
        return {
            "error": str(self),
            "status_code": self.status,
            "source_url": self.source_url,
            "path": self.path,
        }


def cache_dir() -> Path:
    """Return the cache directory, honoring GHED_MCP_CACHE_DIR."""
    raw = os.environ.get("GHED_MCP_CACHE_DIR")
    return Path(raw).expanduser() if raw else Path.home() / ".cache" / "ghed-mcp"


def workbook_path() -> Path:
    return cache_dir() / "ghed.xlsx"


def source_manifest_path() -> Path:
    return cache_dir() / "source-document.json"


def _parse_dotnet_date(value: str | None) -> int:
    if not value:
        return 0
    match = re.search(r"/Date\((-?\d+)\)/", value)
    return int(match.group(1)) if match else 0


def _dotnet_date_to_iso(value: str | None) -> str | None:
    timestamp = _parse_dotnet_date(value)
    if not timestamp:
        return None
    return datetime.fromtimestamp(
        timestamp / 1000, timezone.utc
    ).isoformat().replace("+00:00", "Z")


def _walk_documents(node: dict[str, Any]) -> list[dict[str, Any]]:
    docs = []
    if not node.get("IsFolder"):
        docs.append(node)
    for child in node.get("Children", []) or []:
        docs.extend(_walk_documents(child))
    return docs


def find_latest_all_data_document(tree: dict[str, Any]) -> dict[str, Any]:
    """Find the current GHED all-data workbook in Documentation Centre JSON."""
    candidates = [
        doc for doc in _walk_documents(tree)
        if str(doc.get("Name") or "").lower().startswith("ghed all data")
        and str(doc.get("FileType") or "").lower() == ".xlsx"
        and doc.get("Identifier")
    ]
    if not candidates:
        raise GHEDError("Could not find a public 'GHED all data' workbook.")
    return max(candidates, key=lambda d: _parse_dotnet_date(d.get("DateModified")))


async def get_latest_all_data_document() -> dict[str, Any]:
    """Fetch Documentation Centre metadata and return the current all-data file."""
    last_error: Exception | None = None
    for attempt in range(1, _MAX_DOWNLOAD_ATTEMPTS + 1):
        try:
            async with httpx.AsyncClient(
                timeout=DEFAULT_TIMEOUT,
                headers={"User-Agent": USER_AGENT, "Accept": "application/json"},
                follow_redirects=True,
            ) as client:
                response = await client.get(DOCUMENTATION_TREE_URL)
                response.raise_for_status()
                return find_latest_all_data_document(response.json())
        except httpx.HTTPStatusError as e:
            if not _is_retryable_status(e.response.status_code):
                raise GHEDError(
                    f"HTTP {e.response.status_code} fetching GHED Documentation Centre: "
                    f"{e.response.reason_phrase}",
                    status=e.response.status_code,
                    source_url=DOCUMENTATION_TREE_URL,
                ) from e
            last_error = e
        except httpx.HTTPError as e:
            last_error = e
        if attempt < _MAX_DOWNLOAD_ATTEMPTS:
            delay = _BACKOFF_BASE_SECONDS * (2 ** (attempt - 1))
            logger.warning(
                "Documentation Centre fetch attempt %d failed (%s); retrying in %.0fs",
                attempt,
                last_error,
                delay,
            )
            await asyncio.sleep(delay)
    if isinstance(last_error, httpx.HTTPStatusError):
        raise GHEDError(
            f"HTTP {last_error.response.status_code} fetching GHED Documentation Centre: "
            f"{last_error.response.reason_phrase}",
            status=last_error.response.status_code,
            source_url=DOCUMENTATION_TREE_URL,
        ) from last_error
    raise GHEDError(
        f"Network error fetching GHED Documentation Centre: {last_error}",
        source_url=DOCUMENTATION_TREE_URL,
    ) from last_error


def document_download_url(document_id: int | str) -> str:
    return f"{BASE_URL}/nha/database/DocumentationCentre/GetFile/{document_id}/en"


def normalize_source_document(doc: dict[str, Any]) -> dict[str, Any]:
    document_id = doc.get("Identifier")
    return {
        "document_id": document_id,
        "name": doc.get("Name"),
        "description": doc.get("Description"),
        "file_type": doc.get("FileType"),
        "file_name": doc.get("FileName"),
        "file_size": doc.get("FileSize"),
        "date_modified_raw": doc.get("DateModified"),
        "date_modified": _dotnet_date_to_iso(doc.get("DateModified")),
        "download_url": document_download_url(document_id) if document_id else None,
    }


def read_source_manifest() -> dict[str, Any] | None:
    path = source_manifest_path()
    if not path.exists():
        return None
    try:
        import json

        return json.loads(path.read_text())
    except (OSError, ValueError):
        return None


def write_source_manifest(document: dict[str, Any]) -> None:
    import json

    path = source_manifest_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_document": normalize_source_document(document),
        "downloaded_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
    }
    with tempfile.NamedTemporaryFile(mode="w", dir=path.parent, suffix=".json", delete=False) as fh:
        tmp = Path(fh.name)
        json.dump(payload, fh, indent=2, sort_keys=True)
    try:
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


@asynccontextmanager
async def _download_lock(destination: Path):
    """Cross-process lock, polled without blocking the MCP event loop."""
    lock = FileLock(str(destination) + ".download.lock")
    deadline = time.monotonic() + 600
    while True:
        try:
            lock.acquire(timeout=0)
            break
        except Timeout:
            if time.monotonic() >= deadline:
                raise GHEDError("Timed out waiting for another workbook download.")
            await asyncio.sleep(0.05)
    try:
        yield
    finally:
        lock.release()


def validate_workbook(path: Path) -> None:
    """Reject corrupt archives and incompatible workbook layouts before replacement."""
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ValueError("Archive checksum failed")
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            required = {"Data", "Codebook", "Metadata", "Version"}
            if not required.issubset(wb.sheetnames):
                raise ValueError("Missing required GHED worksheets")
            data = next(wb["Data"].iter_rows(values_only=True))
            if tuple(str(v).lower() for v in data[:5]) != ("location", "code", "region", "income", "year"):
                raise ValueError("Unrecognized Data column layout")
            codebook = next(wb["Codebook"].iter_rows(values_only=True))
            metadata = next(wb["Metadata"].iter_rows(values_only=True))
            if len(codebook) < 8 or str(codebook[0]).lower() != "variable code":
                raise ValueError("Unrecognized Codebook column layout")
            if len(metadata) < 12 or tuple(str(v).lower() for v in metadata[:5]) != (
                "location", "code", "region", "income", "variable code"
            ):
                raise ValueError("Unrecognized Metadata column layout")
        finally:
            wb.close()
    except Exception as exc:
        raise GHEDError(
            f"Invalid GHED workbook: {exc}. Existing cache was preserved.",
            path=str(path),
        ) from exc


async def download_workbook(
    *,
    destination: Path | None = None,
    source_url: str | None = None,
    if_missing: bool = False,
) -> Path:
    """Validate and atomically replace the cache under a cross-process lock."""
    dest = destination or workbook_path()
    dest.parent.mkdir(parents=True, exist_ok=True)
    async with _download_lock(dest):
        if if_missing and dest.exists():
            return dest
        doc = None
        if source_url is None:
            doc = await get_latest_all_data_document()
            source_url = document_download_url(doc["Identifier"])
        with tempfile.NamedTemporaryFile(dir=dest.parent, suffix=".xlsx", delete=False) as fh:
            tmp = Path(fh.name)
        try:
            await _stream_to_file(source_url, tmp)
            await asyncio.to_thread(validate_workbook, tmp)
            tmp.replace(dest)
            if doc is not None:
                write_source_manifest(doc)
            logger.info("Downloaded validated GHED workbook to %s", dest)
            return dest
        except httpx.HTTPStatusError as exc:
            raise GHEDError(
                f"HTTP {exc.response.status_code} downloading GHED workbook",
                status=exc.response.status_code, source_url=source_url, path=str(dest),
            ) from exc
        except (httpx.HTTPError, OSError) as exc:
            raise GHEDError(
                f"Failed to download GHED workbook: {exc}", source_url=source_url, path=str(dest),
            ) from exc
        finally:
            tmp.unlink(missing_ok=True)


async def _stream_to_file(source_url: str, destination: Path) -> None:
    """Stream a GET response to a local file, with retries on transient errors."""
    last_error: Exception | None = None
    for attempt in range(1, _MAX_DOWNLOAD_ATTEMPTS + 1):
        try:
            async with httpx.AsyncClient(
                timeout=DEFAULT_TIMEOUT,
                headers={"User-Agent": USER_AGENT},
                follow_redirects=True,
            ) as client:
                async with client.stream("GET", source_url) as response:
                    response.raise_for_status()
                    with destination.open("wb") as fh:
                        async for chunk in response.aiter_bytes():
                            fh.write(chunk)
            return
        except httpx.HTTPStatusError as e:
            if not _is_retryable_status(e.response.status_code):
                raise
            last_error = e
        except httpx.HTTPError as e:
            last_error = e
        if attempt < _MAX_DOWNLOAD_ATTEMPTS:
            delay = _BACKOFF_BASE_SECONDS * (2 ** (attempt - 1))
            logger.warning(
                "Workbook download attempt %d failed (%s); retrying in %.0fs",
                attempt,
                last_error,
                delay,
            )
            await asyncio.sleep(delay)
    assert last_error is not None
    raise last_error


async def ensure_workbook(refresh: bool = False) -> Path:
    """Return a cached workbook path, downloading it when missing or refreshed."""
    path = workbook_path()
    if refresh or not path.exists():
        return await download_workbook(destination=path, if_missing=not refresh)
    return path


def provenance(
    *,
    workbook: Path | None = None,
    operation: str,
    params: dict[str, Any] | None = None,
    source_signature: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build source metadata for data-returning tool responses."""
    path = workbook or workbook_path()
    retrieved_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    modified_at = None
    size_bytes = None
    if source_signature is not None:
        modified_at = datetime.fromtimestamp(
            source_signature["workbook_mtime_ns"] / 1_000_000_000, timezone.utc
        ).isoformat().replace("+00:00", "Z")
        size_bytes = source_signature["workbook_size_bytes"]
    elif path.exists():
        stat = path.stat()
        modified_at = datetime.fromtimestamp(
            stat.st_mtime, timezone.utc
        ).isoformat().replace("+00:00", "Z")
        size_bytes = stat.st_size
    return {
        "name": "WHO Global Health Expenditure Database",
        "source_url": DOCUMENTATION_TREE_URL,
        "legacy_source_url": LEGACY_SOURCE_URL,
        "workbook_path": str(path),
        "workbook_modified_at": modified_at,
        "workbook_size_bytes": size_bytes,
        **({"dataset_signature": dict(source_signature)} if source_signature is not None else {}),
        "operation": operation,
        "params": dict(params or {}),
        "retrieved_at": retrieved_at,
    }
